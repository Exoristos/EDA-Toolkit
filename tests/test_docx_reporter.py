"""Phase 3a: DOCX reporter tests for eda_toolkit."""

from __future__ import annotations

import os

import pandas as pd
from docx import Document

from eda_toolkit.types import (
    ReportResult,
    PipelineConfig,
    EdaResult,
    ColumnSchema,
    CleaningAction,
    CleaningLevel,
    PiiFinding,
    PiiType,
)
from eda_toolkit.reporters.docx_reporter import render_docx
from eda_toolkit.viz_static import (
    png_correlation_heatmap,
    png_histogram,
    png_missing_bar,
    png_cluster_scatter,
)


# ── Inline fixture helpers (do NOT edit conftest.py) ─────────────────────────

def _make_eda_result(
    df: pd.DataFrame,
    correlation: pd.DataFrame | None = None,
    outliers: dict | None = None,
    hyp_tests: list | None = None,
) -> EdaResult:
    numeric_cols = df.select_dtypes(include=["number"])
    descriptive = numeric_cols.describe() if not numeric_cols.empty else pd.DataFrame()
    missing_summary = pd.DataFrame({
        "column": list(df.columns),
        "missing_count": df.isnull().sum().tolist(),
        "missing_rate": df.isnull().mean().tolist(),
    })
    schemas = [
        ColumnSchema(
            name=col,
            dtype=str(df[col].dtype),
            kind="numeric" if pd.api.types.is_numeric_dtype(df[col]) else "categorical",
            unique_count=int(df[col].nunique()),
            missing_count=int(df[col].isnull().sum()),
        )
        for col in df.columns
    ]
    return EdaResult(
        descriptive=descriptive,
        missing_summary=missing_summary,
        schemas=schemas,
        correlation_matrix=correlation,
        outliers=outliers,
        hypothesis_tests=hyp_tests,
    )


def _make_result(
    df: pd.DataFrame,
    input_path: str = "test.csv",
    pii_findings: list | None = None,
    cleaning_log: list | None = None,
    correlation: pd.DataFrame | None = None,
    outliers: dict | None = None,
    hyp_tests: list | None = None,
) -> ReportResult:
    return ReportResult(
        input_path=input_path,
        row_count=len(df),
        col_count=len(df.columns),
        cleaned_df=df,
        eda=_make_eda_result(df, correlation=correlation, outliers=outliers, hyp_tests=hyp_tests),
        pii_findings=pii_findings if pii_findings is not None else [],
        cleaning_log=cleaning_log if cleaning_log is not None else [],
    )


def _corr_df() -> pd.DataFrame:
    return pd.DataFrame({"a": [1.0, 0.5], "b": [0.5, 1.0]}, index=["a", "b"])


def _pii_findings() -> list[PiiFinding]:
    return [
        PiiFinding(column="tc", pii_type=PiiType.TCKN, confidence=0.95, match_rate=0.80, sample_count=10, detection_layer="text"),
        PiiFinding(column="email", pii_type=PiiType.EMAIL, confidence=0.90, match_rate=0.60, sample_count=5, detection_layer="regex"),
    ]


def _cleaning_log() -> list[CleaningAction]:
    return [
        CleaningAction(column="yas", issue="missing", level=CleaningLevel.CONSERVATIVE, applied=True, detail="filled with median"),
        CleaningAction(column="gelir", issue="outlier", level=CleaningLevel.CONSERVATIVE, applied=False, detail="flagged only"),
    ]


def _outliers_dict() -> dict:
    return {
        "yas": {"count": 3, "rate": 0.015, "lower": 18.0, "upper": 65.0},
        "gelir": {"count": 5, "rate": 0.025, "lower": 1000.0, "upper": 50000.0},
    }


def _hyp_tests() -> list[dict]:
    return [
        {"column": "yas", "test": "shapiro", "statistic": 0.98, "p_value": 0.15, "is_normal": True},
        {"columns": ["yas", "gelir"], "test": "pearson", "statistic": 0.72, "p_value": 0.001, "significant": True},
    ]


# ═══════════════════════════════════════════════════════════════════════════════
# DOCX Reporter Tests
# ═══════════════════════════════════════════════════════════════════════════════


class TestDocxReporter:
    """Tests using synthetic DataFrames with controlled fields."""

    def test_file_created_and_nonempty(self, tmp_path):
        df = pd.DataFrame({"x": [1, 2, 3], "y": [4.0, 5.0, 6.0]})
        result = _make_result(df)
        config = PipelineConfig()
        out = tmp_path / "report.docx"
        render_docx(result, config, str(out))

        assert out.exists()
        assert out.stat().st_size > 0

    def test_creates_parent_dir(self, tmp_path):
        df = pd.DataFrame({"x": [1]})
        result = _make_result(df)
        config = PipelineConfig()
        nested = tmp_path / "deep" / "nested" / "report.docx"
        render_docx(result, config, str(nested))
        assert nested.exists()
        assert nested.stat().st_size > 0

    def test_at_least_one_heading(self, tmp_path):
        """Re-open the docx and assert >= 1 heading."""
        df = pd.DataFrame({"a": [1, 2, 3], "b": [4, 5, 6]})
        result = _make_result(df, pii_findings=_pii_findings(), cleaning_log=_cleaning_log())
        config = PipelineConfig()
        out = tmp_path / "headings.docx"
        render_docx(result, config, str(out))

        doc = Document(str(out))
        headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert len(headings) >= 1, "Expected at least one heading in document"

    def test_embedded_image_when_correlation_present(self, tmp_path):
        """Docx with correlation matrix must contain at least one embedded image."""
        df = pd.DataFrame({"a": [1, 2, 3], "b": [4, 5, 6]})
        result = _make_result(df, correlation=_corr_df())
        config = PipelineConfig()
        out = tmp_path / "with_img.docx"
        render_docx(result, config, str(out))

        doc = Document(str(out))
        # Extract inline shapes (pictures)
        images = []
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                images.append(rel)
        assert len(images) >= 1, f"Expected >=1 embedded image, found {len(images)}"

    def test_full_report_with_all_optional_fields(self, tmp_path):
        """A full report with all optional fields produces a valid larger file."""
        df = pd.DataFrame({"a": [1, 2, 3, 4, 5], "b": [10, 20, 30, 40, 50]})
        result = _make_result(
            df,
            pii_findings=_pii_findings(),
            cleaning_log=_cleaning_log(),
            correlation=_corr_df(),
            outliers=_outliers_dict(),
            hyp_tests=_hyp_tests(),
        )
        config = PipelineConfig()
        out = tmp_path / "full.docx"
        render_docx(result, config, str(out))
        assert out.exists()
        assert out.stat().st_size > 5000, f"Full report too small: {out.stat().st_size} bytes"

    def test_empty_result_does_not_crash(self, tmp_path):
        """Minimal ReportResult with empty data must still produce a valid docx."""
        eda = EdaResult(
            descriptive=pd.DataFrame(),
            missing_summary=pd.DataFrame(),
            schemas=[],
            correlation_matrix=None,
            outliers=None,
            hypothesis_tests=None,
        )
        result = ReportResult(
            input_path="empty.csv",
            row_count=0,
            col_count=0,
            cleaned_df=pd.DataFrame(),
            eda=eda,
            pii_findings=[],
            cleaning_log=[],
        )
        config = PipelineConfig()
        out = tmp_path / "empty.docx"
        render_docx(result, config, str(out))
        assert out.exists()
        assert out.stat().st_size > 0


# ═══════════════════════════════════════════════════════════════════════════════
# viz_static Tests
# ═══════════════════════════════════════════════════════════════════════════════


class TestVizStatic:
    """Ensure each png_* builder returns valid PNG bytes or None."""

    def test_png_correlation_heatmap_nonempty(self):
        corr = _corr_df()
        data = png_correlation_heatmap(corr)
        assert data is not None
        assert isinstance(data, bytes)
        assert data[:4] == b"\x89PNG"
        assert len(data) > 1000

    def test_png_correlation_heatmap_none_on_empty(self):
        assert png_correlation_heatmap(pd.DataFrame()) is None

    def test_png_histogram_nonempty(self):
        series = pd.Series([1.0, 2.0, 3.0, 4.0, 5.0], name="test")
        data = png_histogram(series, "test")
        assert data is not None
        assert data[:4] == b"\x89PNG"
        assert len(data) > 1000

    def test_png_histogram_none_on_empty(self):
        assert png_histogram(pd.Series([], dtype=float), "x") is None

    def test_png_histogram_none_on_all_nan(self):
        series = pd.Series([float("nan"), float("nan")], name="nan")
        assert png_histogram(series, "nan") is None

    def test_png_missing_bar_nonempty(self):
        df = pd.DataFrame({
            "column": ["a", "b", "c"],
            "missing_count": [3, 0, 5],
            "missing_rate": [0.3, 0.0, 0.5],
        })
        data = png_missing_bar(df)
        assert data is not None
        assert data[:4] == b"\x89PNG"

    def test_png_missing_bar_none_on_empty(self):
        assert png_missing_bar(pd.DataFrame()) is None

    def test_png_missing_bar_none_on_no_positive(self):
        df = pd.DataFrame({
            "column": ["x"],
            "missing_count": [0],
            "missing_rate": [0.0],
        })
        assert png_missing_bar(df) is None

    def test_png_missing_bar_none_missing_columns(self):
        df = pd.DataFrame({"foo": [1, 2]})
        assert png_missing_bar(df) is None

    def test_png_cluster_scatter_nonempty(self):
        df = pd.DataFrame({"x": [1, 2, 3, 4, 5], "y": [5, 4, 3, 2, 1]})
        labels = pd.Series([0, 1, 0, 1, 0], name="cluster")
        data = png_cluster_scatter(df, "x", "y", cluster_labels=labels)
        assert data is not None
        assert data[:4] == b"\x89PNG"

    def test_png_cluster_scatter_none_on_empty(self):
        assert png_cluster_scatter(pd.DataFrame(), "x", "y") is None
